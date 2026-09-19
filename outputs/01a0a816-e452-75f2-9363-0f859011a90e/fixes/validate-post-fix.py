import json, hashlib, collections
from pathlib import Path
from openpyxl import load_workbook

fix = Path(__file__).resolve().parent
audit = fix.parent
source = json.loads((audit / 'source-register.json').read_text(encoding='utf-8-sig'))
post = json.loads((fix / 'post-fix-data.json').read_text(encoding='utf-8-sig'))
book_path = fix / 'report/Aven-Post-Fix-Review.xlsx'
book = load_workbook(book_path, data_only=False)
cached = load_workbook(book_path, data_only=True)
assert book.sheetnames == ['Audit', 'Findings', 'Original register']
ids = [f'UX-{n:03d}' for n in range(1, 116)]
assert [book['Audit'].cell(r, 1).value for r in range(13, 128)] == ids
assert [book['Original register'].cell(r, 1).value for r in range(6, 121)] == ids
for n, row in enumerate(post['rows'], 13):
    assert book['Audit'].cell(n, 5).value == row['priorVerdict']
    assert book['Audit'].cell(n, 6).value == row['verdict']
    assert row['finding'] in book['Audit'].cell(n, 10).value
for n, row in enumerate(source['rows'], 6):
    for c, header in enumerate(source['headers'][:16], 1):
        assert (book['Original register'].cell(n, c).value or '') == (row.get(header) or ''), (n, c)
counts = collections.Counter(row['verdict'] for row in post['rows'])
for r in range(5, 11):
    label = book['Audit'].cell(r, 1).value
    assert book['Audit'].cell(r, 2).value.startswith('=COUNTIFS(')
    assert cached['Audit'].cell(r, 2).value == counts[label], (label, cached['Audit'].cell(r, 2).value, counts[label])
assert [book['Findings'].cell(r, 1).value for r in range(11, 16)] == ['UX-058','UX-098','UX-102','UX-066','UX-011']
assert all(book['Findings'].cell(r, 10).value == 'Fixed (tested scope)' for r in range(11, 16))
errors = {'#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A','#NUM!','#NULL!','#SPILL!','#CALC!'}
assert not [(sheet.title,cell.coordinate) for sheet in cached for row in sheet for cell in row if cell.data_type == 'e' or (isinstance(cell.value,str) and cell.value in errors)]
hash_file = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
assert hash_file(source['source']) == source['sha256']
assert hash_file(audit / 'Aven-Independent-Audit.xlsx') == '3fc2940d99c2e99b3d8631c62b9ced8cd89cab0b3a56ca90e7c5925991686d99'
result = {'passed': True, 'sourceRows':115,'sourceCellsCompared':1840,'auditRows':115,'resolvedFeatureRows':5,'distinctDefects':4,'counts':dict(counts),'formulasCachedCorrectly':True,'originalWorkbooksUnchanged':True,'sha256':hash_file(book_path),'bytes':book_path.stat().st_size}
(fix / 'workbook-validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
