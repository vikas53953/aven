import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

root = Path(__file__).parent
register = json.loads((root / 'source-register.json').read_text(encoding='utf-8'))
integrated = json.loads((root / 'integrated-audit.json').read_text(encoding='utf-8'))
book_path = root / 'Aven-Independent-Audit.xlsx'

actual_hash = hashlib.sha256(Path(register['source']).read_bytes()).hexdigest()
assert actual_hash == register['sha256'], (actual_hash, register['sha256'])

wb_formula = load_workbook(book_path, read_only=False, data_only=False)
wb_values = load_workbook(book_path, read_only=True, data_only=True)
assert wb_formula.sheetnames == ['Audit', 'Findings', 'Original register'], wb_formula.sheetnames

audit = wb_formula['Audit']
audit_values = wb_values['Audit']
ids = [audit.cell(row, 1).value for row in range(13, 128)]
expected_ids = [f'UX-{i:03d}' for i in range(1, 116)]
assert ids == expected_ids, (ids[:3], ids[-3:])
assert len(set(ids)) == 115

expected_counts = {'Verified (scoped)': 19, 'Partial': 67, 'Missing': 14, 'Defect': 5, 'Unverified': 6, 'Deferred': 4}
for row in range(5, 11):
    verdict = audit.cell(row, 1).value
    assert audit.cell(row, 2).value.startswith('=COUNTIFS('), audit.cell(row, 2).value
    assert audit_values.cell(row, 2).value == expected_counts[verdict], (verdict, audit_values.cell(row, 2).value)
assert audit.freeze_panes == 'D13', audit.freeze_panes
assert 'Partial can reflect incomplete acceptance coverage' in (audit.cell(2, 1).value or '')
assert '5 affected feature rows; 4 distinct defects' in (audit.cell(10, 5).value or '')

findings = wb_formula['Findings']
finding_values = wb_values['Findings']
assert [findings.cell(10, col).value for col in range(1, 9)] == ['ID', 'Feature', 'Defect title', 'Repro', 'Expected', 'Actual', 'Severity', 'Evidence refs']
finding_ids = [findings.cell(row, 1).value for row in range(11, 16)]
assert finding_ids == ['UX-011', 'UX-066', 'UX-102', 'UX-058', 'UX-098'], finding_ids
assert findings.freeze_panes == 'C11', findings.freeze_panes
assert '5 affected feature rows; 4 distinct defects' in (findings.cell(8, 2).value or '')
assert finding_values.cell(11, 1).value == 'UX-011'

original = wb_formula['Original register']
source_headers = ['ID', 'Area', 'Feature', 'Aven state', 'Aven today', 'Grok Bot', 'Codex desktop', 'Claude Code', 'Aven target / next action', 'Priority', 'Phase', 'Dependency', 'Acceptance check', 'Delivery', 'Decision', 'Evidence']
assert [original.cell(5, col).value for col in range(1, 17)] == source_headers
for offset, source_row in enumerate(register['rows'], start=6):
    actual_row = [original.cell(offset, col).value for col in range(1, 17)]
    expected_row = [source_row.get(header) for header in source_headers]
    assert actual_row == expected_row, (offset, actual_row, expected_row)
assert original.freeze_panes == 'D6', original.freeze_panes

errors = []
for sheet in wb_values.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and any(token in cell.value for token in ['#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!', '#NULL!', '#SPILL!', '#CALC!']):
                errors.append((sheet.title, cell.coordinate, cell.value))
assert not errors, errors[:10]

print(json.dumps({
    'sheets': wb_formula.sheetnames,
    'auditRows': len(ids),
    'uniqueAuditIds': len(set(ids)),
    'formulaCounts': {verdict: audit_values.cell(row, 2).value for row, verdict in zip(range(5, 11), expected_counts)},
    'findingRows': len(finding_ids),
    'originalRows': 115,
    'originalColumns': 16,
    'sourceSha256': actual_hash,
    'formulaErrors': len(errors),
    'freezePanes': {'Audit': audit.freeze_panes, 'Findings': findings.freeze_panes, 'Original register': original.freeze_panes},
}, indent=2))
