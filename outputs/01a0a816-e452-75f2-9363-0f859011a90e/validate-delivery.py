import json
import hashlib
from pathlib import Path
from collections import Counter
import openpyxl

folder = Path(__file__).parent
source = json.loads((folder / 'source-register.json').read_text(encoding='utf-8-sig'))
integrated = json.loads((folder / 'integrated-audit.json').read_text(encoding='utf-8-sig'))
output = folder / 'Aven-Independent-Audit.xlsx'
book = openpyxl.load_workbook(output, data_only=False)
values = openpyxl.load_workbook(output, data_only=True)
assert book.sheetnames == ['Audit', 'Findings', 'Original register'], book.sheetnames
audit = values['Audit']
header_row = next(row[0].row for row in audit.iter_rows() if row[0].value == 'ID' and any(c.value == 'Independent verdict' for c in row))
headers = [c.value for c in audit[header_row]]
columns = {name: i + 1 for i, name in enumerate(headers) if name}
records = {}
for row in audit.iter_rows(min_row=header_row + 1):
    identifier = row[0].value
    if isinstance(identifier, str) and identifier.startswith('UX-'):
        assert identifier not in records, identifier
        records[identifier] = {name: audit.cell(row[0].row, col).value for name, col in columns.items()}
expected = {r['id']: r for r in integrated['rows']}
assert set(records) == set(expected) and len(records) == 115
for identifier, row in expected.items():
    assert records[identifier]['Independent verdict'] == row['verdict'], identifier
    assert records[identifier]['Finding / gap'] == row['finding'], identifier
counts = Counter(r['verdict'] for r in integrated['rows'])
for label, count in counts.items():
    matches = [(cell.row, cell.column) for row in audit.iter_rows(min_row=1, max_row=header_row-1) for cell in row if cell.value == label]
    assert matches, label
    row, col = matches[0]
    assert audit.cell(row, col+1).value == count, (label, audit.cell(row, col+1).value, count)
    assert book['Audit'].cell(row,col+1).data_type == 'f', label
original = values['Original register']
original_header = next(row[0].row for row in original.iter_rows() if row[0].value == 'ID')
for offset, row in enumerate(source['rows'], original_header + 1):
    for col, name in enumerate(source['headers'][:16], 1):
        actual = original.cell(offset, col).value
        expected_value = row.get(name)
        assert (actual or '') == (expected_value or ''), (row['ID'],name,actual,expected_value)
errors = {'#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#N/A', '#NUM!', '#NULL!', '#SPILL!', '#CALC!'}
for sheet in values:
    for row in sheet:
        for cell in row:
            assert not (cell.data_type == 'e' or (isinstance(cell.value,str) and cell.value in errors)), (sheet.title,cell.coordinate,cell.value)
source_hash = hashlib.sha256(Path(source['source']).read_bytes()).hexdigest()
assert source_hash == source['sha256']
baseline = json.loads((folder / 'baseline-hashes.json').read_text(encoding='utf-8-sig'))
changed = [r['path'] for r in baseline if hashlib.sha256(Path(r['path']).read_bytes()).hexdigest() != r['sha256'].lower()]
assert not changed, changed
summary = {'passed':True,'uniqueFeatures':len(records),'counts':dict(counts),'originalSnapshotCellsChecked':115*16,'formulaCountsMatch':True,'sourceWorkbookUnchanged':True,'baselineFilesUnchanged':len(baseline),'outputBytes':output.stat().st_size,'workbookSha256':hashlib.sha256(output.read_bytes()).hexdigest()}
(folder / 'delivery-validation.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
print(json.dumps(summary,indent=2))
